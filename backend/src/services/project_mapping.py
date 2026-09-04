"""项目-城市映射：维护、导入导出与解析。"""

from __future__ import annotations

import io
import secrets
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
import re

from sqlalchemy.ext.asyncio import AsyncSession

from src.db.project_mapping_models import ProjectMappingRecord
from src.models.project_mapping import (
    ProjectMappingCreate,
    ProjectMappingImportResult,
    ProjectMappingPublic,
    ProjectMappingUpdate,
)
from src.repositories.project_mapping import ProjectMappingRepository

EXCEL_HEADERS = (
    "项目名称",
    "项目别名",
    "所在城市",
    "区县",
    "详细地址",
    "标准城市",
    "备注",
)


@dataclass
class ResolvedProjectLocation:
    project_name: str
    city: str
    district: str
    address: str
    policy_city: str
    full_location: str


def _new_id() -> str:
    return f"proj_{secrets.token_hex(4)}"


def _to_public(record: ProjectMappingRecord) -> ProjectMappingPublic:
    return ProjectMappingPublic(
        id=record.id,
        project_name=record.project_name,
        aliases=record.aliases or "",
        city=record.city or "",
        district=record.district or "",
        address=record.address or "",
        policy_city=record.policy_city or "",
        remark=record.remark or "",
        updated_by=record.updated_by or "",
        updated_at=record.updated_at.isoformat() if record.updated_at else "",
    )


def _split_aliases(raw: str) -> list[str]:
    if not raw:
        return []
    return [part.strip() for part in raw.replace("，", ",").split(",") if part.strip()]


def _full_location(city: str, district: str, address: str) -> str:
    parts: list[str] = []
    if city:
        parts.append(city if city.endswith("市") else f"{city}市")
    if district:
        parts.append(district)
    if address:
        parts.append(address)
    return "".join(parts)


class ProjectMappingService:
    def __init__(self, db: AsyncSession):
        self.repo = ProjectMappingRepository(db)

    async def list_mappings(
        self, keyword: str | None = None
    ) -> list[ProjectMappingPublic]:
        records = await self.repo.list_all(keyword)
        return [_to_public(r) for r in records]

    async def create(
        self, body: ProjectMappingCreate, updated_by: str
    ) -> ProjectMappingPublic:
        existing = await self.repo.get_by_name(body.project_name.strip())
        if existing:
            raise ValueError(f"项目名称已存在：{body.project_name}")
        record = ProjectMappingRecord(
            id=_new_id(),
            project_name=body.project_name.strip(),
            aliases=body.aliases.strip(),
            city=body.city.strip(),
            district=body.district.strip(),
            address=body.address.strip(),
            policy_city=(body.policy_city or body.city).strip(),
            remark=body.remark.strip(),
            updated_by=updated_by,
        )
        created = await self.repo.create(record)
        return _to_public(created)

    async def update(
        self, mapping_id: str, body: ProjectMappingUpdate, updated_by: str
    ) -> ProjectMappingPublic | None:
        record = await self.repo.get_by_id(mapping_id)
        if record is None:
            return None
        fields = body.model_dump(exclude_unset=True)
        if "project_name" in fields and fields["project_name"]:
            fields["project_name"] = fields["project_name"].strip()
        if "policy_city" in fields and fields["policy_city"]:
            fields["policy_city"] = fields["policy_city"].strip()
        fields["updated_by"] = updated_by
        updated = await self.repo.update(record, **fields)
        return _to_public(updated)

    async def delete(self, mapping_id: str) -> bool:
        record = await self.repo.get_by_id(mapping_id)
        if record is None:
            return False
        await self.repo.delete(record)
        return True

    async def resolve_from_text(self, text: str) -> ResolvedProjectLocation | None:
        if not text.strip():
            return None
        records = await self.repo.list_all()
        best: tuple[int, ProjectMappingRecord] | None = None

        def consider(score: int, record: ProjectMappingRecord) -> None:
            nonlocal best
            if best is None or score > best[0]:
                best = (score, record)

        for record in records:
            candidates = [record.project_name, *_split_aliases(record.aliases)]
            for name in candidates:
                if not name:
                    continue
                if name in text:
                    consider(len(name) + 100, record)

        if best is None:
            hints = re.findall(r"[\u4e00-\u9fffA-Za-z0-9·]{2,24}", text)
            for hint in sorted(set(hints), key=len, reverse=True):
                for record in records:
                    candidates = [record.project_name, *_split_aliases(record.aliases)]
                    for name in candidates:
                        if not name or len(hint) < 2:
                            continue
                        if name.startswith(hint) or hint.startswith(name):
                            consider(len(hint) + (50 if name.startswith(hint) else 10), record)

        if best is None:
            return None
        record = best[1]
        city = record.city or ""
        district = record.district or ""
        address = record.address or ""
        policy_city = record.policy_city or city
        return ResolvedProjectLocation(
            project_name=record.project_name,
            city=city,
            district=district,
            address=address,
            policy_city=policy_city,
            full_location=_full_location(city, district, address),
        )

    async def normalize_memory_structured(
        self,
        structured: dict,
        *,
        source_text: str = "",
    ) -> dict:
        from src.agent.memory_normalizer import normalize_structured_fields

        records = await self.repo.list_all()
        return normalize_structured_fields(structured, records, source_text=source_text)

    async def import_excel(
        self, content: bytes, updated_by: str, *, replace: bool = False
    ) -> ProjectMappingImportResult:
        if replace:
            await self.repo.delete_all()

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return ProjectMappingImportResult(imported=0, updated=0, skipped=0, errors=["空文件"])

        header = [str(c or "").strip() for c in rows[0]]
        if header[: len(EXCEL_HEADERS)] != list(EXCEL_HEADERS):
            return ProjectMappingImportResult(
                imported=0,
                updated=0,
                skipped=0,
                errors=[f"表头须为：{','.join(EXCEL_HEADERS)}"],
            )

        imported = updated = skipped = 0
        errors: list[str] = []

        for idx, row in enumerate(rows[1:], start=2):
            cells = [str(c or "").strip() if c is not None else "" for c in row]
            while len(cells) < len(EXCEL_HEADERS):
                cells.append("")
            project_name = cells[0]
            if not project_name:
                skipped += 1
                continue
            payload = {
                "aliases": cells[1],
                "city": cells[2],
                "district": cells[3],
                "address": cells[4],
                "policy_city": cells[5] or cells[2],
                "remark": cells[6],
                "updated_by": updated_by,
            }
            existing = await self.repo.get_by_name(project_name)
            try:
                if existing:
                    await self.repo.update(existing, **payload)
                    updated += 1
                else:
                    record = ProjectMappingRecord(
                        id=_new_id(),
                        project_name=project_name,
                        **payload,
                    )
                    await self.repo.create(record)
                    imported += 1
            except Exception as exc:  # noqa: BLE001
                errors.append(f"第{idx}行：{project_name} — {exc}")

        return ProjectMappingImportResult(
            imported=imported, updated=updated, skipped=skipped, errors=errors[:20]
        )

    def build_template_excel(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "项目城市映射"
        ws.append(list(EXCEL_HEADERS))
        ws.append(
            [
                "燕宝能源",
                "雁宝,燕宝可视化二期",
                "呼伦贝尔",
                "海拉尔区",
                "燕宝能源大厦",
                "海拉尔",
                "驻场办公点",
            ]
        )
        ws.append(
            [
                "神东能源数据治理平台",
                "神东项目",
                "鄂尔多斯",
                "",
                "神东能源大厦",
                "鄂尔多斯",
                "",
            ]
        )
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    async def export_excel(self) -> bytes:
        records = await self.repo.list_all()
        wb = Workbook()
        ws = wb.active
        ws.title = "项目城市映射"
        ws.append(list(EXCEL_HEADERS))
        for record in records:
            ws.append(
                [
                    record.project_name,
                    record.aliases,
                    record.city,
                    record.district,
                    record.address,
                    record.policy_city,
                    record.remark,
                ]
            )
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
